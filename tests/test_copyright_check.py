import json
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import copyright_check
from copyright_check import Findings, _contains_material_value, check_materials, check_sources, redact_line


class RedactLineTests(unittest.TestCase):
    """脱敏只替换密钥本身，不得破坏标识符、赋值号和语法结构。"""

    def test_assignment_keeps_identifier_and_syntax(self):
        cases = [
            ('const API_KEY = "sk-live-4b8f2a91c3d7e5064b8f";',
             'const API_KEY = "****";'),
            ('spring.datasource.password = "Abc123456!"',
             'spring.datasource.password = "****"'),
            ('private static final String APP_SECRET = "9f8e7d6c5b4a39281706";',
             'private static final String APP_SECRET = "****";'),
            ("access_token: 'ya29.a0AfH6SMBxxxxxxxxxxxx'",
             "access_token: '****'"),
            ('appsecret = "0123456789abcdef0123456789abcdef"',
             'appsecret = "****"'),
        ]
        for source, expected in cases:
            with self.subTest(source=source):
                self.assertEqual(redact_line(source), expected)

    def test_bare_token_keeps_scheme_prefix(self):
        """裸令牌保留 4 位方案前缀，便于说明这里原本是什么类型的凭据。"""
        self.assertEqual(redact_line('AWS_KEY = "AKIAIOSFODNN7EXAMPLE"'),
                         'AWS_KEY = "AKIA****"')
        self.assertEqual(redact_line('curl -H "Authorization: ghp_1234567890abcdefghijklmnopqrstuvwxyz"'),
                         'curl -H "Authorization: ghp_****"')

    def test_password_value_is_not_partially_leaked(self):
        """口令本身不保留任何明文前缀。"""
        self.assertNotIn("Abc1", redact_line('password = "Abc123456!"'))

    def test_redaction_is_idempotent(self):
        once = redact_line('const API_KEY = "sk-live-4b8f2a91c3d7e5064b8f";')
        self.assertEqual(redact_line(once), once)

    def test_business_code_is_untouched(self):
        line = "  return items.reduce((s, i) => s + i.price, 0);"
        self.assertEqual(redact_line(line), line)

    def test_pii_is_still_masked(self):
        self.assertEqual(redact_line('const PHONE = "13800138000";'),
                         'const PHONE = "138******00";')
        self.assertIn("***@corp-intranet.cn", redact_line("contact: dev@corp-intranet.cn"))

    def test_allowlisted_pii_is_kept(self):
        line = "import type { Foo } from '@types/foo'"
        self.assertEqual(redact_line(line), line)


class SecretFindingTests(unittest.TestCase):
    """同一处密钥只应报一条高风险，避免闸门计数虚高。"""

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.repo = Path(self.temp.name)

    def _scan(self, source):
        (self.repo / "app.js").write_text(source, encoding="utf-8")
        findings = Findings()
        check_sources(self.repo, {"copyright_holder": "Demo Co.,Ltd.",
                                  "projects": [{"id": "01-demo", "name": "demo",
                                                "source_files": ["app.js"]}]}, findings)
        return [i for i in findings.items if i["rule"].startswith("敏感信息")]

    def test_one_secret_reports_one_finding(self):
        hits = self._scan('const API_KEY = "sk-live-4b8f2a91c3d7e5064b8f";\n')
        self.assertEqual(len(hits), 1, f"同一处密钥被重复计数：{[h['rule'] for h in hits]}")

    def test_two_distinct_secrets_on_one_line_report_two(self):
        hits = self._scan('run("AKIAIOSFODNN7EXAMPLE", "ghp_1234567890abcdefghijklmnopqrstuvwxyz")\n')
        self.assertEqual(len(hits), 2)

    def test_underscore_prefixed_names_are_detected(self):
        """.env / docker-compose 里最常见的写法，不能因为 \\b 被下划线挡住而漏检。"""
        for source in ['const DB_PASSWORD = "Abc123456!";',
                       'MYSQL_ROOT_PASSWORD = "root1234"',
                       'REDIS_PASSWORD: "abcdef"',
                       'WECHAT_APP_SECRET = "0123456789abcdef"']:
            with self.subTest(source=source):
                self.assertEqual(len(self._scan(source + "\n")), 1)

    def test_similar_names_are_not_flagged(self):
        self.assertEqual(self._scan('const passwordHash = "notasecret";\n'), [])

    def test_reported_detail_is_masked(self):
        hits = self._scan('const API_KEY = "sk-live-4b8f2a91c3d7e5064b8f";\n')
        self.assertNotIn("4b8f2a91c3d7e506", hits[0]["detail"])

    def test_report_does_not_leak_password_prefix(self):
        """报告是本地留档文件，不应写入口令的任何明文片段。"""
        hits = self._scan('const DB_PASSWORD = "Abc123456!";\n')
        self.assertNotIn("Abc1", hits[0]["detail"])


class MaterialRoleTests(unittest.TestCase):
    def test_pdf_text_spacing_does_not_break_name_match(self):
        self.assertTrue(_contains_material_value("益搭 AI 协作系统（服务端）", "益搭AI协作系统（服务端）"))

    def test_new_source_filename_is_not_checked_as_manual(self):
        """“源码”交付件是程序材料，不应因不含软件全称被误报。"""
        from docx import Document
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            project = {"id": "01-demo", "name": "演示软件（服务端）", "version": "V1.0",
                       "short_name": "演示", "artifact_label": "后端", "source_pages": 60}
            d = root / project["id"] / "源程序提取"
            d.mkdir(parents=True)
            doc = Document()
            doc.add_paragraph("演示-后端源码")
            doc.save(d / "演示-后端源码.docx")
            findings = Findings()
            check_materials(root, {"copyright_holder": "Demo", "projects": [project]}, findings)
            self.assertFalse(any(item["rule"] == "材料中找不到软件全称" for item in findings.items))

    def test_manifest_limits_material_scan_to_formal_files(self):
        """旧版目录中的仓库地址不应污染正式材料检查。"""
        from docx import Document
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            project = {"id": "01-demo", "name": "演示软件", "version": "V1.0",
                       "short_name": "演示", "artifact_label": "后端", "source_pages": 60}
            project_dir = root / project["id"]
            source_dir = project_dir / "源程序提取"
            source_dir.mkdir(parents=True)
            for path, title in [
                (source_dir / "演示-后端源码.docx", "演示源码"),
                (project_dir / "演示-后端软件说明.docx", "演示软件 V1.0"),
            ]:
                if path.suffix == ".docx":
                    doc = Document(); doc.add_paragraph(title); doc.save(path)
                else:
                    path.write_bytes(b"%PDF-1.4")
            old = project_dir / "旧版未裁剪材料"; old.mkdir()
            doc = Document(); doc.add_paragraph("https://github.com/example/old"); doc.save(old / "旧.docx")
            (project_dir / "材料上传清单.json").write_text(json.dumps({
                "_schema": "ruanzhu-kit.material-manifest.v1",
                "materials": {
                    "programPdf": {"path": "01-demo/源程序提取/演示-后端源码.docx"},
                    "docPdf": {"path": "01-demo/演示-后端软件说明.docx"},
                },
            }, ensure_ascii=False), encoding="utf-8")
            findings = Findings()
            check_materials(root, {"copyright_holder": "Demo", "projects": [project]}, findings)
            self.assertFalse(any("旧版未裁剪材料" in item["where"] for item in findings.items))


if __name__ == "__main__":
    unittest.main()
