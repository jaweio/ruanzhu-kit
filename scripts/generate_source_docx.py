#!/usr/bin/env python3

import argparse
import json
import sys
from pathlib import Path
from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from oss_scrub import own_tokens  # noqa: E402
from source_material_render import build_source_html, collect_source_material, read_lines  # noqa: E402
from output_names import source_material_docx_name, source_material_html_name  # noqa: E402


def set_cell_no_wrap(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_pr.append(OxmlElement("w:noWrap"))


def set_cell_margins(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin in ("top", "left", "bottom", "right"):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), "0")
        node.set(qn("w:type"), "dxa")


def remove_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = borders.find(qn(f"w:{edge}"))
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            borders.append(elem)
        elem.set(qn("w:val"), "nil")


def fixed_layout(table):
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def add_page(doc, page_lines, lines_per_page, font_size_pt, row_height_pt):
    table = doc.add_table(rows=lines_per_page, cols=1)
    table.autofit = False
    remove_borders(table)
    fixed_layout(table)
    for i, row in enumerate(table.rows):
        row.height = Pt(row_height_pt)
        row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
        cell = row.cells[0]
        cell.width = Cm(19)
        set_cell_no_wrap(cell)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(page_lines[i] or " ")
        run.font.size = Pt(font_size_pt)


def build_doc(repo, out_root, project, tokens=(), cli_keep_comments=False, cli_keep_imports=False):
    lines_per_page = int(project.get("lines_per_page", 90))
    pages = int(project.get("source_pages", 60))
    font_size = float(project.get("source_font_size", 6.5))
    row_height = float(project.get("source_row_height", 8.15))
    collected = collect_source_material(repo, project, tokens, cli_keep_comments, cli_keep_imports)
    lines = collected["lines"]
    selected = collected["selected"]
    missing = collected["missing"]
    skipped = collected["skipped"]
    scrub_stats = collected["stats"]
    redact = project.get("redact", True)
    scrub = project.get("scrub_open_source", True)
    trim_comments = collected["trim_comments"]
    trim_imports = collected["trim_imports"]
    max_blank_lines = collected["max_blank_lines"]
    need = pages * lines_per_page
    if skipped:
        print(f"[{project['name']}] 跳过 {len(skipped)} 个第三方文件：" + "、".join(r for r, _ in skipped), file=sys.stderr)
    if len(lines) < need:
        print(f"[{project['name']}] 警告：可用代码 {len(lines)} 行，不足 {need} 行（{pages} 页），请补充自研文件",
              file=sys.stderr)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(0.65)
    section.bottom_margin = Cm(0.65)
    section.left_margin = Cm(1)
    section.right_margin = Cm(1)

    for page in range(pages):
        chunk = selected[page * lines_per_page : (page + 1) * lines_per_page]
        add_page(doc, chunk, lines_per_page, font_size, row_height)
        if page < pages - 1:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    out_dir = out_root / project["id"] / "源程序提取"
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / source_material_docx_name(project)
    doc.save(out_path)
    html_path = out_dir / source_material_html_name(project)
    html_path.write_text(
        # 共享同一批 selected 行，确保 DOCX 与 Chromium PDF 内容一致。
        build_source_html(project, selected, pages, lines_per_page),
        encoding="utf-8",
    )
    (out_dir / "源程序DOCX生成报告.md").write_text(
        f"# {project['name']} 源程序 DOCX 生成报告\n\n"
        f"- 输出文件：`{out_path.name}`\n"
        f"- Chromium PDF 输入：`{html_path.name}`\n"
        f"- 页面方向：竖向 A4\n"
        f"- 页数：{pages}\n"
        f"- 每页代码行数：{lines_per_page}\n"
        f"- 缺失文件：{'、'.join(missing) if missing else '无'}\n"
        f"- 敏感信息脱敏：{'已开启（密钥/手机号/邮箱/内网 IP）' if redact else '未开启'}\n"
        f"- 开源痕迹清除：{'已开启' if scrub else '未开启'}（删除自有许可证/版权头 {scrub_stats['headers']} 处，"
        f"含仓库地址或开源字样的注释行 {scrub_stats['dropped']} 行）\n"
        f"- 材料压缩：普通注释 {'已移除' if trim_comments else '保留'} {scrub_stats['comments']} 行；"
        f"导入/include/use 声明 {'已移除' if trim_imports else '保留'} {scrub_stats['imports']} 行；"
        f"连续空行已压缩（最多 {max_blank_lines} 行）\n"
        f"- 自动跳过的第三方文件：{len(skipped)} 个\n"
        + "".join(f"  - `{r}`：{why}\n" for r, why in skipped)
        + f"- 可用代码行数：{len(lines)}（需要 {need}）\n",
        encoding="utf-8",
    )
    print(out_path)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", default="soft-copyright-materials/ruanzhu.config.json")
    parser.add_argument("--repo", default=".")
    parser.add_argument("--keep-comments", action="store_true", help="保留普通代码注释（默认裁剪）")
    parser.add_argument("--keep-imports", action="store_true", help="保留 import/include/use 声明（默认裁剪）")
    parser.add_argument("--preview", action="store_true", help="同时生成源程序材料 HTML 可视化预览")
    parser.add_argument("--preview-out", help="预览 HTML 输出路径")
    parser.add_argument("--preview-json", help="同时输出预览统计 JSON")
    parser.add_argument("--preview-sample-lines", type=int, default=120, help="每个文件预览最多展示多少行")
    args = parser.parse_args()
    cfg = json.loads(Path(args.config).read_text(encoding="utf-8"))
    out_root = Path(cfg.get("output_root", "soft-copyright-materials"))
    repo = Path(args.repo)
    tokens = own_tokens(cfg, repo)
    for project in cfg["projects"]:
        build_doc(repo, out_root, project, tokens, args.keep_comments, args.keep_imports)
    if args.preview:
        from source_preview import write_preview
        preview_out = Path(args.preview_out) if args.preview_out else out_root / "源程序材料预览.html"
        write_preview(
            cfg, Path(args.config), repo, preview_out,
            False if args.keep_comments else None,
            False if args.keep_imports else None,
            None, max(1, args.preview_sample_lines), args.preview_json,
        )
        print(preview_out.resolve())


if __name__ == "__main__":
    main()
